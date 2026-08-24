"""Import monthly tabbed commission-details Excel → draft Commission Invoices.

Expected layout (one sheet per product), header row with columns like:
  SR.NO | SAIL MONTH | INVOICE NO. | VESSEL SAIL DATE | PO. NO. | PO DATE |
  SHIP TO PARTY | PORT OF DISCHARGE | QTY | CURRENCY | RATE PER MT | VALUE |
  OCEAN FREIGHT | INSURANCE | FOB VALUE | COMMISSION | CURRENCY

One data row → one draft CI (GBInc). Bill-to always GBL; bank Chemung defaults.
"""
from __future__ import annotations

import json
import re
import uuid
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.commission_invoices import (
    VARIANT_GBINC,
    _float,
    calculate_ci_totals,
    create_commission_invoice,
    dollars_in_words,
    get_default_ci,
    invoice_date_from_sail,
    recalc_line,
)
from app.database import get_data_dir

# Normalized header → logical key
_HEADER_ALIASES: dict[str, str] = {
    "sr.no": "sr_no",
    "sr no": "sr_no",
    "srno": "sr_no",
    "sail month": "sail_month",
    "invoice no.": "invoice_no",
    "invoice no": "invoice_no",
    "invoice number": "invoice_no",
    "vessel sail date": "sail_date",
    "sail date": "sail_date",
    "po. no.": "po_no",
    "po. no": "po_no",
    "po no.": "po_no",
    "po no": "po_no",
    "po number": "po_no",
    "po date": "po_date",
    "ship to party": "ship_to",
    "ship to": "ship_to",
    "port of discharge": "port",
    "qty": "qty",
    "quantity": "qty",
    "currency": "currency",
    "rate per mt": "rate",
    "rate": "rate",
    "value": "value",
    "ocean frieght": "freight",  # common typo in sheets
    "ocean freight": "freight",
    "freight": "freight",
    "insurance": "insurance",
    "fob value": "fob",
    "fob": "fob",
    "commission": "commission",
}


def _norm_header(val: Any) -> str:
    s = re.sub(r"\s+", " ", str(val or "").strip().lower())
    return s


def _parse_date(val: Any) -> str:
    """Parse spreadsheet dates as India format (day-month-year). Return YYYY-MM-DD or ''."""
    if val is None or val == "":
        return ""
    if isinstance(val, datetime):
        return val.date().isoformat()
    if isinstance(val, date):
        return val.isoformat()

    s = str(val).strip()
    # Drop trailing notes after the date token
    s = re.split(r"\s{2,}|\s+\(", s, maxsplit=1)[0].strip()
    # Normalize typos: 03.05..2026 → 03.05.2026 ; mixed separators
    s = re.sub(r"[./\-]+", ".", s)

    # ISO already unambiguous
    m_iso = re.match(r"^(\d{4})\.(\d{1,2})\.(\d{1,2})", s)
    if m_iso:
        y, mo, d = int(m_iso.group(1)), int(m_iso.group(2)), int(m_iso.group(3))
        try:
            return date(y, mo, d).isoformat()
        except ValueError:
            return ""

    # India / GBL sheets: day.month.year (also dd/mm/yyyy, dd-mm-yyyy)
    m = re.match(r"^(\d{1,2})\.(\d{1,2})\.(\d{2,4})", s)
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if y < 100:
            y += 2000
        try:
            return date(y, mo, d).isoformat()
        except ValueError:
            return ""

    # Text months — still day-first where possible
    for fmt in ("%d-%b-%Y", "%d %b %Y", "%d-%B-%Y", "%d %B %Y"):
        try:
            return datetime.strptime(s[:20].title(), fmt).date().isoformat()
        except ValueError:
            continue
    return ""


def _cell_num(val: Any) -> float:
    if val is None or val == "":
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    return _float(str(val).replace(",", ""))


def _product_from_sheet(ws, sheet_name: str) -> str:
    title = str(ws.cell(2, 2).value or ws.cell(1, 1).value or "").strip()
    m = re.search(
        r"SUPPLY\s+OF\s+(.+)$",
        title,
        re.IGNORECASE,
    )
    if m:
        return m.group(1).strip()
    return (sheet_name or "").strip()


def _find_header_row(ws) -> tuple[int, dict[str, int]]:
    """Return (row_index, {logical_key: col_index})."""
    for r in range(1, min(ws.max_row, 15) + 1):
        mapping: dict[str, int] = {}
        for c in range(1, min(ws.max_column, 30) + 1):
            key = _HEADER_ALIASES.get(_norm_header(ws.cell(r, c).value))
            if key and key not in mapping:
                # First "currency" → currency; second → commission_currency (ignore)
                if key == "currency" and "currency" in mapping:
                    continue
                mapping[key] = c
        if "invoice_no" in mapping and ("qty" in mapping or "rate" in mapping):
            return r, mapping
    return 0, {}


def _row_dict(ws, r: int, colmap: dict[str, int]) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for key, c in colmap.items():
        out[key] = ws.cell(r, c).value
    return out


def _is_total_row(row: dict[str, Any]) -> bool:
    for key in ("sr_no", "invoice_no", "ship_to", "sail_month"):
        v = row.get(key)
        if v is not None and str(v).strip().upper() == "TOTAL":
            return True
    return False


def _po_additional_info(po_no: str, po_date: str, sail_month: str) -> str:
    parts = []
    if po_no:
        parts.append(f"PO No: {po_no}")
    if po_date:
        parts.append(f"PO Date: {po_date}")
    if sail_month:
        parts.append(f"Sail month: {sail_month}")
    return " | ".join(parts)


def _build_ci_from_row(
    row: dict[str, Any],
    *,
    product: str,
    sheet: str,
    variant: str = VARIANT_GBINC,
) -> dict[str, Any]:
    inv_no = str(row.get("invoice_no") or "").strip()
    ship_to = str(row.get("ship_to") or "").strip()
    port = str(row.get("port") or "").strip()
    po_no = str(row.get("po_no") or "").strip()
    sail_month = str(row.get("sail_month") or "").strip()
    sail_date = _parse_date(row.get("sail_date"))
    po_date = _parse_date(row.get("po_date"))
    # If PO date parse failed, keep raw short text
    if not po_date and row.get("po_date") not in (None, ""):
        po_date = str(row.get("po_date")).strip()[:32]

    qty = _cell_num(row.get("qty"))
    rate = _cell_num(row.get("rate"))
    freight = _cell_num(row.get("freight"))
    insurance = _cell_num(row.get("insurance"))
    value = _cell_num(row.get("value"))
    if not value and qty and rate:
        value = round(qty * rate, 2)
    fob = _cell_num(row.get("fob"))
    if not fob and value:
        fob = round(value - freight - insurance, 2)
    elif not fob and qty and rate:
        fob = round(qty * rate - freight - insurance, 2)

    cur = str(row.get("currency") or "USD").strip().upper() or "USD"
    if cur not in ("USD", "EUR", "INR", "GBP"):
        cur = "USD"

    commission_rate = 3.0
    # If sheet has explicit commission number (not formula result empty), keep 3% default
    # unless we can infer rate from commission/fob
    comm_cell = _cell_num(row.get("commission"))
    if fob and comm_cell:
        inferred = round(comm_cell / fob * 100, 4)
        if 0.1 < inferred < 50:
            commission_rate = inferred

    ci = get_default_ci(variant)
    # Bill-to stays Godavari Biorefineries Ltd (defaults)
    ci["invoice_number"] = inv_no
    ci["shipment_date"] = sail_date
    ci["invoice_date"] = invoice_date_from_sail(sail_date) or date.today().isoformat()
    ci["notice_date"] = ci["invoice_date"]
    ci["customer_order_no"] = po_no
    ci["delivery_port"] = port
    ci["fob_currency"] = cur
    ci["value_currency"] = cur
    ci["status"] = "Draft"
    ci["transaction_description"] = (
        f"Commission for supply of {product}" + (f" to {ship_to}" if ship_to else "")
    )
    additional = _po_additional_info(po_no, po_date, sail_month)
    ci["enclosures"] = additional
    if freight or insurance:
        note = f"Imported from {sheet}. Ocean freight: {freight:g}; Insurance: {insurance:g}."
        ci["internal_notes"] = note
    else:
        ci["internal_notes"] = f"Imported from sheet {sheet}."

    line = {
        "end_customer": ship_to,
        "product_description": product,
        "gbl_invoice_number": inv_no,
        "quantity": qty,
        # Keep unit_price empty so the editor does not overwrite freight-adjusted FOB
        # as qty × rate. CIF holds rate for the Notice of Order page.
        "unit_price": 0.0,
        "cif_price": rate,
        "fob_value": fob,
        "commission_rate": commission_rate,
        "commission_value": 0.0,
        "shipment_date": sail_date,
    }
    line = recalc_line(line)
    totals = calculate_ci_totals([line], 0)
    ci["line_items"] = totals["line_items"]
    ci["total_commission"] = totals["total_commission"]
    ci["amount_in_words"] = dollars_in_words(totals["total_commission"], cur)

    missing: dict[str, str] = {}
    if not inv_no:
        missing["invoice_number"] = "missing"
    if not ship_to:
        missing["end_customer"] = "missing"
    if not product:
        missing["product_description"] = "missing"
    if not qty:
        missing["quantity"] = "missing"
    if not fob:
        missing["fob_value"] = "missing"
    if not port:
        missing["delivery_port"] = "missing"
    if not sail_date:
        missing["shipment_date"] = "missing"
    if not po_no:
        missing["customer_order_no"] = "missing"
    ci["_field_hints"] = missing

    ci["_import_meta"] = {
        "sheet": sheet,
        "product": product,
        "sail_month": sail_month,
        "po_date": po_date,
        "freight": freight,
        "insurance": insurance,
        "value": value,
    }
    return ci


def parse_commission_details_workbook(
    data: bytes,
    *,
    variant: str = VARIANT_GBINC,
) -> list[dict[str, Any]]:
    """Parse uploaded .xlsx bytes → list of draft CI dicts (not yet saved)."""
    from io import BytesIO

    wb = load_workbook(BytesIO(data), data_only=True)
    drafts: list[dict[str, Any]] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        header_row, colmap = _find_header_row(ws)
        if not header_row or not colmap:
            continue
        product = _product_from_sheet(ws, sheet_name)
        for r in range(header_row + 1, ws.max_row + 1):
            row = _row_dict(ws, r, colmap)
            if _is_total_row(row):
                continue
            inv = str(row.get("invoice_no") or "").strip()
            qty = _cell_num(row.get("qty"))
            ship = str(row.get("ship_to") or "").strip()
            if not inv and not qty and not ship:
                continue
            drafts.append(
                _build_ci_from_row(row, product=product, sheet=sheet_name, variant=variant)
            )
    return drafts


def preview_rows(drafts: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Compact rows for the preview table."""
    existing = existing_ci_invoice_keys()
    out = []
    for i, ci in enumerate(drafts):
        line = (ci.get("line_items") or [{}])[0]
        meta = ci.get("_import_meta") or {}
        inv = ci.get("invoice_number") or ""
        already = _normalize_invoice_key(inv) in existing if inv else False
        out.append({
            "index": i,
            "sheet": meta.get("sheet") or "",
            "product": line.get("product_description") or "",
            "invoice_number": inv,
            "ship_to": line.get("end_customer") or "",
            "port": ci.get("delivery_port") or "",
            "qty": line.get("quantity") or 0,
            "fob": line.get("fob_value") or 0,
            "commission": line.get("commission_value") or 0,
            "currency": ci.get("value_currency") or "USD",
            "shipment_date": ci.get("shipment_date") or "",
            "po_no": ci.get("customer_order_no") or "",
            "missing": sorted((ci.get("_field_hints") or {}).keys()),
            "already_exists": already,
            "selected_default": not already,
        })
    return out


def _import_dir() -> Path:
    d = get_data_dir() / "tmp" / "ci_excel_import"
    d.mkdir(parents=True, exist_ok=True)
    return d


def save_import_batch(drafts: list[dict[str, Any]]) -> str:
    token = uuid.uuid4().hex
    path = _import_dir() / f"{token}.json"
    path.write_text(json.dumps(drafts, default=str), encoding="utf-8")
    return token


def load_import_batch(token: str) -> list[dict[str, Any]] | None:
    if not token or not re.fullmatch(r"[a-f0-9]{32}", token):
        return None
    path = _import_dir() / f"{token}.json"
    if not path.is_file():
        return None
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None
    return data if isinstance(data, list) else None


def clear_import_batch(token: str) -> None:
    if not token or not re.fullmatch(r"[a-f0-9]{32}", token):
        return
    path = _import_dir() / f"{token}.json"
    try:
        path.unlink(missing_ok=True)
    except OSError:
        pass


def create_drafts_from_batch(
    token: str,
    indices: list[int],
    *,
    variant: str = VARIANT_GBINC,
) -> list[int]:
    """Persist selected draft CIs; return new CI ids."""
    drafts = load_import_batch(token)
    if not drafts:
        return []
    created: list[int] = []
    for i in indices:
        if i < 0 or i >= len(drafts):
            continue
        ci = dict(drafts[i])
        ci.pop("_import_meta", None)
        hints = ci.pop("_field_hints", None)
        line_items = ci.pop("line_items", []) or []
        ci["variant"] = variant
        ci["status"] = "Draft"
        # Keep hints in internal_notes only if useful — store as JSON? skip
        if hints and not ci.get("prepared_by"):
            pass
        ci_id = create_commission_invoice(ci, line_items, variant=variant)
        created.append(ci_id)
    clear_import_batch(token)
    return created


def _normalize_invoice_key(invoice: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", (invoice or "").strip().upper())


def existing_ci_invoice_keys(*, variant: str = VARIANT_GBINC) -> set[str]:
    """Normalized invoice numbers already on non-cancelled CIs."""
    try:
        from app.commission_invoices import upgrade_commission_invoices_schema
        from app.database import get_db

        upgrade_commission_invoices_schema()
        with get_db() as conn:
            rows = conn.execute(
                """
                SELECT invoice_number FROM commission_invoices
                WHERE COALESCE(variant, 'gbinc') = ?
                  AND LOWER(COALESCE(status, 'Draft')) != 'cancelled'
                """,
                (variant,),
            ).fetchall()
    except Exception:
        return set()
    return {
        _normalize_invoice_key(r["invoice_number"] if hasattr(r, "keys") else r[0])
        for r in rows
        if (r["invoice_number"] if hasattr(r, "keys") else r[0])
    }


def drafts_from_register_rows(
    rows: list[dict[str, Any]],
    *,
    variant: str = VARIANT_GBINC,
) -> list[dict[str, Any]]:
    """Map GBL CS register rows → draft CI dicts (same shape as Excel import)."""
    drafts: list[dict[str, Any]] = []
    for row in rows:
        inv = str(row.get("invoice_no") or "").strip()
        qty = _cell_num(row.get("qty"))
        ship = str(row.get("ship_to_party") or "").strip()
        if not inv and not qty and not ship:
            continue
        product = str(row.get("product_hint") or "").strip() or "Product"
        sheet = str(row.get("source_sheet") or row.get("source_original") or "register")
        mapped = {
            "invoice_no": inv,
            "ship_to": ship,
            "port": row.get("port_of_discharge"),
            "po_no": row.get("po_no"),
            "po_date": row.get("po_date"),
            "sail_month": row.get("sail_month"),
            "sail_date": row.get("vessel_sail_date"),
            "qty": row.get("qty"),
            "rate": row.get("rate_per_mt"),
            "freight": row.get("ocean_freight"),
            "insurance": row.get("insurance"),
            "value": row.get("value"),
            "fob": row.get("fob_value"),
            "currency": row.get("currency") or row.get("commission_currency") or "USD",
            "commission": row.get("commission"),
        }
        ci = _build_ci_from_row(mapped, product=product, sheet=sheet, variant=variant)
        meta = ci.setdefault("_import_meta", {})
        meta["register_id"] = row.get("id")
        meta["source_file"] = row.get("source_original") or row.get("source_file") or ""
        drafts.append(ci)
    return drafts
