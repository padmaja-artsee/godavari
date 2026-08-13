"""Excel and PDF exports for Actual vs Projection."""
from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

try:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    _HAS_REPORTLAB = True
except ImportError:  # pragma: no cover
    colors = TA_CENTER = TA_LEFT = TA_RIGHT = None  # type: ignore
    landscape = letter = ParagraphStyle = getSampleStyleSheet = None  # type: ignore
    inch = Paragraph = SimpleDocTemplate = Spacer = Table = TableStyle = None  # type: ignore
    _HAS_REPORTLAB = False

try:
    from PIL import Image, ImageDraw, ImageFont
except ImportError:  # pragma: no cover
    Image = ImageDraw = ImageFont = None  # type: ignore

_GREEN = "1C5631"
_GREEN_D = "134023"
_GREEN_L = "E8F5E9"
_WHITE = "FFFFFF"
_BLACK = "1A1A2E"
_GREY_B = "E0E3E8"


def _font(bold=False, size=9, color=_BLACK):
    return Font(name="Calibri", size=size, bold=bold, color=color)


def _fill(hex_color: str):
    return PatternFill("solid", fgColor=hex_color)


def _border():
    s = Side(style="thin", color=_GREY_B)
    return Border(left=s, right=s, top=s, bottom=s)


def _al(h="right", v="center", wrap=True):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


def _fmt_num(value, decimals=None) -> str | float | int:
    if value is None or value == "":
        return ""
    try:
        n = float(value)
    except (TypeError, ValueError):
        return value
    if decimals == 0:
        return round(n)
    if decimals is not None:
        return round(n, decimals)
    if n == int(n):
        return int(n)
    return round(n, 3)


def _fmt_money(value) -> str | float | int:
    if value is None or value == "":
        return ""
    try:
        return round(float(value))
    except (TypeError, ValueError):
        return value


def export_avp_xlsx(report: dict[str, Any]) -> tuple[bytes, str]:
    wb = Workbook()
    ws = wb.active
    fy = report["fiscal_year"]
    month = report["month_label"]
    ws.title = f"AVP {report['fy_label']} {month}"[:31]

    # Title
    ws.merge_cells("A1:H1")
    c = ws["A1"]
    c.value = "GBInc — Actual vs Projection"
    c.font = _font(bold=True, size=13, color=_WHITE)
    c.fill = _fill(_GREEN)
    c.alignment = _al(h="left")
    ws.row_dimensions[1].height = 22

    ws.merge_cells("A2:H2")
    c = ws["A2"]
    c.value = f"{report['fy_label']}  ·  Month: {report['month'].title()}"
    c.font = _font(bold=True, size=10, color=_WHITE)
    c.fill = _fill(_GREEN_D)
    c.alignment = _al(h="left")

    # Group headers
    hdr1 = 4
    ws.cell(row=hdr1, column=1, value="Product").font = _font(bold=True, size=9, color=_WHITE)
    ws.cell(row=hdr1, column=1).fill = _fill(_GREEN_D)
    ws.cell(row=hdr1, column=1).alignment = _al(h="center")
    ws.merge_cells(start_row=hdr1, start_column=1, end_row=hdr1 + 1, end_column=1)

    groups = [
        (2, 3, report["month_header"]),
        (4, 5, report["ytd_header"]),
        (6, 8, report["target_header"]),
    ]
    for start, end, label in groups:
        ws.merge_cells(start_row=hdr1, start_column=start, end_row=hdr1, end_column=end)
        cell = ws.cell(row=hdr1, column=start, value=label)
        cell.font = _font(bold=True, size=9, color=_WHITE)
        cell.fill = _fill(_GREEN)
        cell.alignment = _al(h="center")
        for col in range(start, end + 1):
            ws.cell(row=hdr1, column=col).fill = _fill(_GREEN)
            ws.cell(row=hdr1, column=col).border = _border()

    # Sub headers
    hdr2 = 5
    sub = [
        (2, "Qty in MT"),
        (3, "Income in $"),
        (4, "Qty in MT"),
        (5, "Income in $"),
        (6, "Qty in MT"),
        (7, "Sales Val $"),
        (8, "Income in $"),
    ]
    for col, label in sub:
        cell = ws.cell(row=hdr2, column=col, value=label)
        cell.font = _font(bold=True, size=8, color=_WHITE)
        cell.fill = _fill(_GREEN_D)
        cell.alignment = _al(h="center")
        cell.border = _border()
    ws.cell(row=hdr2, column=1).border = _border()
    ws.cell(row=hdr2, column=1).fill = _fill(_GREEN_D)

    # Data rows
    r = 6
    for row in report["rows"]:
        values = [
            row["product"],
            _fmt_num(row["month_qty"]),
            _fmt_money(row["month_income"]),
            _fmt_num(row["ytd_qty"]),
            _fmt_money(row["ytd_income"]),
            _fmt_num(row["target_qty"]),
            _fmt_money(row["target_sales"]),
            _fmt_money(row["target_income"]),
        ]
        for col, val in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=val)
            cell.font = _font(size=9)
            cell.border = _border()
            cell.alignment = _al(h="left" if col == 1 else "right")
            if col > 1 and isinstance(val, (int, float)):
                cell.number_format = "#,##0.###" if col in (2, 4, 6) else "#,##0"
        r += 1

    # Totals
    tot = report["totals"]
    totals = [
        "Total",
        _fmt_num(tot["month_qty"]),
        _fmt_money(tot["month_income"]),
        _fmt_num(tot["ytd_qty"]),
        _fmt_money(tot["ytd_income"]),
        _fmt_num(tot["target_qty"]),
        _fmt_money(tot["target_sales"]),
        _fmt_money(tot["target_income"]),
    ]
    for col, val in enumerate(totals, start=1):
        cell = ws.cell(row=r, column=col, value=val)
        cell.font = _font(bold=True, size=9)
        cell.fill = _fill(_GREEN_L)
        cell.border = _border()
        cell.alignment = _al(h="left" if col == 1 else "right")
        if col > 1 and isinstance(val, (int, float)):
            cell.number_format = "#,##0.###" if col in (2, 4, 6) else "#,##0"

    widths = [22, 12, 12, 12, 12, 12, 14, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = BytesIO()
    wb.save(buf)
    fname = f"GBInc_Actual_vs_Projection_{report['fy_label'].replace(' ', '')}_{month}.xlsx"
    return buf.getvalue(), fname


def export_avp_pdf(report: dict[str, Any]) -> tuple[bytes, str]:
    if not _HAS_REPORTLAB:
        raise RuntimeError("reportlab is required for PDF export — pip install reportlab")
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(letter),
        leftMargin=0.5 * inch,
        rightMargin=0.5 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "AVPTitle",
        parent=styles["Heading1"],
        fontSize=14,
        textColor=colors.HexColor("#1C5631"),
        spaceAfter=4,
        alignment=TA_LEFT,
    )
    sub_style = ParagraphStyle(
        "AVPSub",
        parent=styles["Normal"],
        fontSize=9,
        textColor=colors.HexColor("#6B7280"),
        spaceAfter=12,
    )
    cell_left = ParagraphStyle("AVPL", parent=styles["Normal"], fontSize=8, leading=10)
    cell_right = ParagraphStyle(
        "AVPR", parent=styles["Normal"], fontSize=8, leading=10, alignment=TA_RIGHT
    )
    hdr_style = ParagraphStyle(
        "AVPH",
        parent=styles["Normal"],
        fontSize=7,
        leading=9,
        textColor=colors.white,
        alignment=TA_CENTER,
    )

    def H(text: str) -> Paragraph:
        return Paragraph(text.replace("\n", "<br/>"), hdr_style)

    def L(text: str) -> Paragraph:
        return Paragraph(str(text), cell_left)

    def R(text: str) -> Paragraph:
        return Paragraph(str(text), cell_right)

    def money(v) -> str:
        if v is None or v == "":
            return ""
        try:
            return f"{float(v):,.0f}"
        except (TypeError, ValueError):
            return str(v)

    def qty(v) -> str:
        if v is None or v == "":
            return ""
        try:
            n = float(v)
            return f"{n:g}" if n == int(n) else f"{n:,.3f}".rstrip("0").rstrip(".")
        except (TypeError, ValueError):
            return str(v)

    story = [
        Paragraph("GBInc — Actual vs Projection", title_style),
        Paragraph(
            f"{report['fy_label']} · Month: {report['month'].title()}",
            sub_style,
        ),
    ]

    header1 = [
        H("Product"),
        H(report["month_header"]),
        "",
        H(report["ytd_header"]),
        "",
        H(report["target_header"]),
        "",
        "",
    ]
    header2 = [
        H(""),
        H("Qty<br/>in MT"),
        H("Income<br/>in $"),
        H("Qty<br/>in MT"),
        H("Income<br/>in $"),
        H("Qty<br/>in MT"),
        H("Sales Val<br/>$"),
        H("Income<br/>in $"),
    ]
    data = [header1, header2]
    for row in report["rows"]:
        data.append(
            [
                L(row["product"]),
                R(qty(row["month_qty"])),
                R(money(row["month_income"])),
                R(qty(row["ytd_qty"])),
                R(money(row["ytd_income"])),
                R(qty(row["target_qty"])),
                R(money(row["target_sales"])),
                R(money(row["target_income"])),
            ]
        )
    tot = report["totals"]
    data.append(
        [
            L("<b>Total</b>"),
            R(f"<b>{qty(tot['month_qty'])}</b>"),
            R(f"<b>{money(tot['month_income'])}</b>"),
            R(f"<b>{qty(tot['ytd_qty'])}</b>"),
            R(f"<b>{money(tot['ytd_income'])}</b>"),
            R(f"<b>{qty(tot['target_qty'])}</b>"),
            R(f"<b>{money(tot['target_sales'])}</b>"),
            R(f"<b>{money(tot['target_income'])}</b>"),
        ]
    )

    col_widths = [1.6 * inch, 0.9 * inch, 0.95 * inch, 0.9 * inch, 0.95 * inch, 0.9 * inch, 1.05 * inch, 0.95 * inch]
    table = Table(data, colWidths=col_widths, repeatRows=2)
    green = colors.HexColor("#1C5631")
    green_d = colors.HexColor("#134023")
    green_l = colors.HexColor("#E8F5E9")
    style_cmds = [
        ("BACKGROUND", (0, 0), (-1, 1), green_d),
        ("TEXTCOLOR", (0, 0), (-1, 1), colors.white),
        ("BACKGROUND", (1, 0), (2, 0), green),
        ("BACKGROUND", (3, 0), (4, 0), green),
        ("BACKGROUND", (5, 0), (7, 0), green),
        ("SPAN", (1, 0), (2, 0)),
        ("SPAN", (3, 0), (4, 0)),
        ("SPAN", (5, 0), (7, 0)),
        ("SPAN", (0, 0), (0, 1)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (1, 2), (-1, -1), "RIGHT"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
        ("BACKGROUND", (0, -1), (-1, -1), green_l),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]
    # zebra rows
    for i in range(2, len(data) - 1):
        if i % 2 == 0:
            style_cmds.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#FAFBFC")))
    table.setStyle(TableStyle(style_cmds))
    story.append(table)
    story.append(Spacer(1, 0.2 * inch))

    doc.build(story)
    fname = f"GBInc_Actual_vs_Projection_{report['fy_label'].replace(' ', '')}_{report['month_label']}.pdf"
    return buf.getvalue(), fname


def _png_font(size: int, bold: bool = False):
    """Best-effort system font; falls back to Pillow default."""
    if ImageFont is None:
        return None
    candidates = []
    if bold:
        candidates.extend(
            [
                "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
                "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
                "/Library/Fonts/Arial Bold.ttf",
                "/System/Library/Fonts/Helvetica.ttc",
            ]
        )
    candidates.extend(
        [
            "/System/Library/Fonts/Supplemental/Arial.ttf",
            "/Library/Fonts/Arial.ttf",
            "/System/Library/Fonts/Helvetica.ttc",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
        ]
    )
    for path in candidates:
        try:
            return ImageFont.truetype(path, size=size)
        except Exception:
            continue
    return ImageFont.load_default()


def _png_text_w(draw, text: str, font) -> int:
    if hasattr(draw, "textbbox"):
        box = draw.textbbox((0, 0), text, font=font)
        return box[2] - box[0]
    return len(text) * 7


def export_avp_png(report: dict[str, Any]) -> tuple[bytes, str]:
    """Render Actual vs Projection as a PNG image (Pillow)."""
    if Image is None:
        raise RuntimeError("Pillow is required for PNG export")

    def money(v) -> str:
        if v is None or v == "":
            return ""
        try:
            return f"{float(v):,.0f}"
        except (TypeError, ValueError):
            return str(v)

    def qty(v) -> str:
        if v is None or v == "":
            return ""
        try:
            n = float(v)
            return f"{n:g}" if n == int(n) else f"{n:,.3f}".rstrip("0").rstrip(".")
        except (TypeError, ValueError):
            return str(v)

    headers_sub = [
        "",
        "Qty in MT",
        "Income in $",
        "Qty in MT",
        "Income in $",
        "Qty in MT",
        "Sales Val $",
        "Income in $",
    ]

    body: list[list[str]] = []
    for row in report["rows"]:
        body.append(
            [
                str(row["product"]),
                qty(row["month_qty"]),
                money(row["month_income"]),
                qty(row["ytd_qty"]),
                money(row["ytd_income"]),
                qty(row["target_qty"]),
                money(row["target_sales"]),
                money(row["target_income"]),
            ]
        )
    tot = report["totals"]
    body.append(
        [
            "Total",
            qty(tot["month_qty"]),
            money(tot["month_income"]),
            qty(tot["ytd_qty"]),
            money(tot["ytd_income"]),
            qty(tot["target_qty"]),
            money(tot["target_sales"]),
            money(tot["target_income"]),
        ]
    )

    font_title = _png_font(22, bold=True)
    font_sub = _png_font(13)
    font_hdr = _png_font(11, bold=True)
    font_cell = _png_font(12)
    font_cell_b = _png_font(12, bold=True)

    pad_x = 14
    # Measure with a scratch image
    scratch = Image.new("RGB", (10, 10), "white")
    d0 = ImageDraw.Draw(scratch)

    col_mins = [160, 90, 100, 90, 100, 90, 110, 100]
    col_w = list(col_mins)
    for row in body:
        for i, cell in enumerate(row):
            col_w[i] = max(col_w[i], _png_text_w(d0, cell, font_cell) + pad_x * 2)
    for i, h in enumerate(headers_sub):
        if h:
            col_w[i] = max(col_w[i], _png_text_w(d0, h, font_hdr) + pad_x * 2)
    # Group header widths
    col_w[0] = max(col_w[0], _png_text_w(d0, "Product", font_hdr) + pad_x * 2)

    table_w = sum(col_w)
    row_h = 34
    header_h = 32
    title_h = 70
    margin = 28
    img_w = table_w + margin * 2
    img_h = title_h + header_h * 2 + row_h * len(body) + margin * 2

    green = (28, 86, 49)
    green_d = (19, 64, 35)
    green_l = (232, 245, 233)
    white = (255, 255, 255)
    text = (26, 26, 46)
    muted = (107, 114, 128)
    zebra = (250, 251, 252)
    grid = (203, 213, 225)

    img = Image.new("RGB", (img_w, img_h), white)
    draw = ImageDraw.Draw(img)

    x0, y0 = margin, margin
    draw.text((x0, y0), "GBInc — Actual vs Projection", fill=green, font=font_title)
    draw.text(
        (x0, y0 + 30),
        f"{report['fy_label']} · Month: {report['month'].title()}",
        fill=muted,
        font=font_sub,
    )

    y = y0 + title_h
    # Header row 1 (grouped)
    spans = [(0, 0), (1, 2), (3, 4), (5, 7)]
    labels = ["Product", report["month_header"], report["ytd_header"], report["target_header"]]
    xs = [x0]
    for w in col_w[:-1]:
        xs.append(xs[-1] + w)

    draw.rectangle([x0, y, x0 + table_w, y + header_h], fill=green_d)
    for (c0, c1), label in zip(spans, labels):
        left = xs[c0]
        right = xs[c1] + col_w[c1]
        if c0 > 0:
            draw.rectangle([left, y, right, y + header_h], fill=green)
        tw = _png_text_w(draw, label, font_hdr)
        draw.text(
            (left + (right - left - tw) / 2, y + (header_h - 14) / 2),
            label,
            fill=white,
            font=font_hdr,
        )
    y += header_h

    # Header row 2
    draw.rectangle([x0, y, x0 + table_w, y + header_h], fill=green_d)
    for i, h in enumerate(headers_sub):
        if not h:
            continue
        left = xs[i]
        right = left + col_w[i]
        tw = _png_text_w(draw, h, font_hdr)
        draw.text(
            (left + (right - left - tw) / 2, y + (header_h - 14) / 2),
            h,
            fill=white,
            font=font_hdr,
        )
    y += header_h

    for ri, row in enumerate(body):
        is_total = ri == len(body) - 1
        bg = green_l if is_total else (zebra if ri % 2 == 0 else white)
        draw.rectangle([x0, y, x0 + table_w, y + row_h], fill=bg)
        font = font_cell_b if is_total else font_cell
        for i, cell in enumerate(row):
            left = xs[i]
            right = left + col_w[i]
            tw = _png_text_w(draw, cell, font)
            if i == 0:
                tx = left + pad_x
            else:
                tx = right - pad_x - tw
            draw.text((tx, y + (row_h - 14) / 2), cell, fill=text, font=font)
        y += row_h

    # Grid lines
    y_top = y0 + title_h
    y_bot = y
    draw.rectangle([x0, y_top, x0 + table_w, y_bot], outline=grid, width=1)
    for i in range(1, 8):
        draw.line([(xs[i], y_top), (xs[i], y_bot)], fill=grid, width=1)
    yy = y_top
    for i in range(2 + len(body) + 1):
        draw.line([(x0, yy), (x0 + table_w, yy)], fill=grid, width=1)
        if i < 2:
            yy += header_h
        elif i < 2 + len(body):
            yy += row_h

    out = BytesIO()
    img.save(out, format="PNG", optimize=True)
    fname = f"GBInc_Actual_vs_Projection_{report['fy_label'].replace(' ', '')}_{report['month_label']}.png"
    return out.getvalue(), fname
