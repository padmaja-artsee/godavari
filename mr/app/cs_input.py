"""Parse GBL commission statement (CS) Excel uploads and keep a cumulative register."""
from __future__ import annotations

import re
import calendar
import sqlite3
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.database import get_data_dir

MONTH_NAMES = {
    1: "JANUARY", 2: "FEBRUARY", 3: "MARCH", 4: "APRIL",
    5: "MAY", 6: "JUNE", 7: "JULY", 8: "AUGUST",
    9: "SEPTEMBER", 10: "OCTOBER", 11: "NOVEMBER", 12: "DECEMBER",
}

# Display order matches the consolidated management-report spreadsheet.
DISPLAY_COLUMNS = [
    ("sr_no", "SR.NO"),
    ("sail_month", "SAIL MONTH"),
    ("invoice_no", "INVOICE NO."),
    ("vessel_sail_date", "VESSEL SAIL DATE"),
    ("po_no", "PO. NO."),
    ("po_date", "PO DATE"),
    ("ship_to_party", "SHIP TO PARTY"),
    ("port_of_discharge", "PORT OF DISCHARGE"),
    ("qty", "QTY"),
    ("currency", "CURRENCY"),
    ("rate_per_mt", "RATE PER MT"),
    ("value", "VALUE"),
    ("ocean_freight", "OCEAN FREIGHT"),
    ("insurance", "INSURANCE"),
    ("fob_value", "FOB VALUE"),
    ("commission", "COMMISSION"),
    ("commission_currency", "CURRENCY"),
]

# Header aliases → canonical field key (handles spacing/typos/optional PO cols).
HEADER_ALIASES = {
    "SR.NO": "sr_no",
    "SR NO": "sr_no",
    "SRNO": "sr_no",
    "SAIL MONTH": "sail_month",
    "INVOICE NO.": "invoice_no",
    "INVOICE NO": "invoice_no",
    "INVOICE NUMBER": "invoice_no",
    "VESSEL SAIL DATE": "vessel_sail_date",
    "PO. NO.": "po_no",
    "PO. NO": "po_no",
    "PO NO.": "po_no",
    "PO NO": "po_no",
    "PO NUMBER": "po_no",
    "PO DATE": "po_date",
    "SHIP TO PARTY": "ship_to_party",
    "PORT OF DISCHARGE": "port_of_discharge",
    "QTY": "qty",
    "QUANTITY": "qty",
    "CURRENCY": "currency",  # first occurrence; second remapped below
    "RATE PER MT": "rate_per_mt",
    "VALUE": "value",
    "OCEAN FREIGHT": "ocean_freight",
    "OCEAN FRIEGHT": "ocean_freight",
    "INSURANCE": "insurance",
    "FOB VALUE": "fob_value",
    "COMMISSION": "commission",
}

TEXT_FIELDS = {
    "sr_no", "sail_month", "invoice_no", "vessel_sail_date", "po_no", "po_date",
    "ship_to_party", "port_of_discharge", "currency", "commission_currency",
}
NUMERIC_FIELDS = {
    "qty", "rate_per_mt", "value", "ocean_freight", "insurance", "fob_value", "commission",
}


def cs_upload_dir() -> Path:
    d = get_data_dir() / "uploads" / "mr" / "gbl_cs"
    d.mkdir(parents=True, exist_ok=True)
    return d


def register_db_path() -> Path:
    d = get_data_dir() / "mr"
    d.mkdir(parents=True, exist_ok=True)
    return d / "gbl_cs_register.db"


def _norm(text: Any) -> str:
    if text is None:
        return ""
    s = str(text).strip().upper()
    s = re.sub(r"\s+", " ", s)
    return s


def _cell_str(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def _to_number(v: Any):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


_MONTH_NAME_TO_NUM = {
    "JAN": 1, "JANUARY": 1,
    "FEB": 2, "FEBRUARY": 2,
    "MAR": 3, "MARCH": 3,
    "APR": 4, "APRIL": 4,
    "MAY": 5,
    "JUN": 6, "JUNE": 6,
    "JUL": 7, "JULY": 7,
    "AUG": 8, "AUGUST": 8,
    "SEP": 9, "SEPT": 9, "SEPTEMBER": 9,
    "OCT": 10, "OCTOBER": 10,
    "NOV": 11, "NOVEMBER": 11,
    "DEC": 12, "DECEMBER": 12,
}


def _clamp_date(year: int, month: int, day: int) -> datetime | None:
    """Build a date, clamping day to the last valid day of the month.

    Spreadsheets sometimes use impossible days (e.g. 'June 31 2026').
    """
    if year < 100:
        year += 2000
    if month < 1 or month > 12:
        return None
    last = calendar.monthrange(year, month)[1]
    day = max(1, min(int(day), last))
    try:
        return datetime(year, month, day)
    except ValueError:
        return None


def _parse_sail_date(value: Any):
    """Parse vessel sail date; returns datetime or None.

    Handles DD.MM.YYYY, double-dot typos (03.05..2026), month-name forms
    like 'June 31 2026', and trailing notes like '31.07.2026 (TODAY SAILING…)'.
    Invalid days are clamped to the last day of that month.
    """
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    # Excel serial dates sometimes come through as float
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            from openpyxl.utils.datetime import from_excel
            return from_excel(value)
        except Exception:
            return None

    s = str(value).strip()
    # Numeric forms: DD.MM.YYYY (also double-dot typos)
    m = re.match(r"^(\d{1,2})[./-](\d{1,2})[./-]+(\d{2,4})", s)
    if m:
        day, month, year = int(m.group(1)), int(m.group(2)), int(m.group(3))
        return _clamp_date(year, month, day)

    # Month-name forms: "June 31 2026", "31 June 2026", "Jun-30-2026"
    month_names = (
        "January|February|March|April|May|June|July|August|September|"
        "October|November|December|Jan|Feb|Mar|Apr|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec"
    )
    m = re.match(
        rf"^({month_names})[.\-/\s]+(\d{{1,2}})(?:st|nd|rd|th)?[.\-/\s,]+(\d{{2,4}})",
        s,
        re.I,
    )
    if m:
        mon = _MONTH_NAME_TO_NUM.get(m.group(1).upper())
        if mon:
            return _clamp_date(int(m.group(3)), mon, int(m.group(2)))
    m = re.match(
        rf"^(\d{{1,2}})(?:st|nd|rd|th)?[.\-/\s]+({month_names})[.\-/\s,]+(\d{{2,4}})",
        s,
        re.I,
    )
    if m:
        mon = _MONTH_NAME_TO_NUM.get(m.group(2).upper())
        if mon:
            return _clamp_date(int(m.group(3)), mon, int(m.group(1)))
    return None


def _month_from_sail_date(value: Any) -> str:
    """Canonical sail month from vessel sail date (preferred over sheet SAIL MONTH)."""
    dt = _parse_sail_date(value)
    if not dt:
        return ""
    return MONTH_NAMES.get(dt.month, "")


def _looks_like_header(cells: list[Any]) -> bool:
    norms = {_norm(c) for c in cells if c is not None}
    return (
        "SR.NO" in norms
        or "SR NO" in norms
        or any(n.startswith("INVOICE") for n in norms)
    )


def _is_total_row(cells: list[Any]) -> bool:
    return any(_norm(c) == "TOTAL" for c in cells if c is not None)


def _map_headers(header_cells: list[Any]) -> dict[int, str]:
    """Map 0-based column index → field key. Second CURRENCY → commission_currency."""
    mapping: dict[int, str] = {}
    seen_currency = False
    for i, cell in enumerate(header_cells):
        key = HEADER_ALIASES.get(_norm(cell))
        if not key:
            continue
        if key == "currency":
            if seen_currency:
                mapping[i] = "commission_currency"
            else:
                mapping[i] = "currency"
                seen_currency = True
        else:
            mapping[i] = key
    return mapping


def _product_from_title(title: str) -> str:
    m = re.search(r"SUPPLY OF\s+(.+)$", title or "", re.I)
    return (m.group(1).strip() if m else "") or ""


def _parse_sheet(ws, sheet_name: str) -> dict[str, Any] | None:
    title = ""
    header_row = None
    start_col = 1
    header_cells: list[Any] = []

    for r in range(1, min(ws.max_row, 40) + 1):
        row_vals = [ws.cell(r, c).value for c in range(1, (ws.max_column or 1) + 1)]
        non_empty = [v for v in row_vals if v is not None and str(v).strip()]
        if not non_empty:
            continue
        if not title and len(non_empty) == 1 and "COMMISSION" in _norm(non_empty[0]):
            title = str(non_empty[0]).strip()
            continue
        if _looks_like_header(row_vals):
            header_row = r
            for c, v in enumerate(row_vals, start=1):
                if v is not None and str(v).strip():
                    start_col = c
                    break
            header_cells = row_vals[start_col - 1 :]
            break

    if header_row is None:
        return None

    col_map = _map_headers(header_cells)
    if "invoice_no" not in col_map.values() and "sail_month" not in col_map.values():
        return None

    rows: list[dict[str, Any]] = []
    for r in range(header_row + 1, (ws.max_row or 0) + 1):
        cells = [ws.cell(r, start_col + i).value for i in range(len(header_cells))]
        if all(c is None or str(c).strip() == "" for c in cells):
            continue
        if _is_total_row(cells):
            break

        record = {key: None for key, _ in DISPLAY_COLUMNS}
        for i, field in col_map.items():
            if i >= len(cells):
                continue
            val = cells[i]
            if field in TEXT_FIELDS:
                record[field] = _cell_str(val)
            elif field in NUMERIC_FIELDS:
                record[field] = _to_number(val)
            else:
                record[field] = val

        # Require at least an invoice or a ship-to party to count as a line
        if not record.get("invoice_no") and not record.get("ship_to_party"):
            continue
        # Sail month is driven by vessel sail date (not the sheet's SAIL MONTH label)
        month_from_date = _month_from_sail_date(record.get("vessel_sail_date"))
        if month_from_date:
            record["sail_month"] = month_from_date
        elif record.get("sail_month"):
            record["sail_month"] = str(record["sail_month"]).strip().upper()
        record["source_sheet"] = sheet_name
        record["product_hint"] = _product_from_title(title) or sheet_name
        rows.append(record)

    if not rows:
        return None

    return {
        "sheet_name": sheet_name,
        "title": title or f"GBL Commission Statement ({sheet_name})",
        "rows": rows,
        "row_count": len(rows),
    }


def parse_gbl_cs_xlsx(data: bytes) -> dict[str, Any]:
    """Parse all commission sheets in a workbook into a unified row list."""
    wb = load_workbook(BytesIO(data), data_only=True)
    sheets = []
    all_rows: list[dict[str, Any]] = []

    for name in wb.sheetnames:
        parsed = _parse_sheet(wb[name], name)
        if not parsed:
            continue
        sheets.append({"sheet_name": parsed["sheet_name"], "title": parsed["title"],
                       "row_count": parsed["row_count"]})
        all_rows.extend(parsed["rows"])

    if not all_rows:
        raise ValueError(
            "Could not find commission statement rows "
            "(expected columns like SR.NO, INVOICE NO., COMMISSION)."
        )

    return {
        "sheets": sheets,
        "title": sheets[0]["title"] if len(sheets) == 1 else f"{len(sheets)} sheets",
        "sheet_name": ", ".join(s["sheet_name"] for s in sheets),
        "column_labels": [label for _, label in DISPLAY_COLUMNS],
        "rows": all_rows,
        "row_count": len(all_rows),
        "totals": None,
    }


def save_cs_upload(filename: str, data: bytes) -> Path:
    safe = re.sub(r"[^\w.\- ]+", "_", Path(filename).name) or "gbl_cs.xlsx"
    dest = cs_upload_dir() / safe
    if dest.exists():
        stem, suffix = dest.stem, dest.suffix
        n = 2
        while True:
            cand = dest.with_name(f"{stem}_{n}{suffix}")
            if not cand.exists():
                dest = cand
                break
            n += 1
    dest.write_bytes(data)
    return dest


# ── Cumulative register (SQLite) ─────────────────────────────────────────────

def _connect() -> sqlite3.Connection:
    conn = sqlite3.connect(str(register_db_path()), timeout=30)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode=WAL")
    return conn


def init_register() -> None:
    with _connect() as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS gbl_cs_lines (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                invoice_no TEXT NOT NULL,
                sr_no TEXT,
                sail_month TEXT,
                vessel_sail_date TEXT,
                po_no TEXT,
                po_date TEXT,
                ship_to_party TEXT,
                port_of_discharge TEXT,
                qty REAL,
                currency TEXT,
                rate_per_mt REAL,
                value REAL,
                ocean_freight REAL,
                insurance REAL,
                fob_value REAL,
                commission REAL,
                commission_currency TEXT,
                source_file TEXT,
                source_original TEXT,
                source_sheet TEXT,
                product_hint TEXT,
                uploaded_at TEXT NOT NULL
            )
            """
        )
        # Migrate older DBs that enforced UNIQUE(invoice_no) — recreate without it.
        cols = {r[1] for r in conn.execute("PRAGMA table_info(gbl_cs_lines)").fetchall()}
        create_sql = conn.execute(
            "SELECT sql FROM sqlite_master WHERE type='table' AND name='gbl_cs_lines'"
        ).fetchone()
        needs_rebuild = False
        if create_sql and create_sql[0] and "invoice_no TEXT NOT NULL UNIQUE" in create_sql[0]:
            needs_rebuild = True
        if "source_original" not in cols:
            needs_rebuild = True
        if needs_rebuild:
            conn.execute(
                """
                CREATE TABLE gbl_cs_lines_new (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    invoice_no TEXT NOT NULL,
                    sr_no TEXT,
                    sail_month TEXT,
                    vessel_sail_date TEXT,
                    po_no TEXT,
                    po_date TEXT,
                    ship_to_party TEXT,
                    port_of_discharge TEXT,
                    qty REAL,
                    currency TEXT,
                    rate_per_mt REAL,
                    value REAL,
                    ocean_freight REAL,
                    insurance REAL,
                    fob_value REAL,
                    commission REAL,
                    commission_currency TEXT,
                    source_file TEXT,
                    source_original TEXT,
                    source_sheet TEXT,
                    product_hint TEXT,
                    uploaded_at TEXT NOT NULL
                )
                """
            )
            common = [
                "invoice_no", "sr_no", "sail_month", "vessel_sail_date", "po_no", "po_date",
                "ship_to_party", "port_of_discharge", "qty", "currency", "rate_per_mt",
                "value", "ocean_freight", "insurance", "fob_value", "commission",
                "commission_currency", "source_file", "source_sheet", "product_hint",
                "uploaded_at",
            ]
            have = [c for c in common if c in cols]
            select_cols = ", ".join(have)
            insert_cols = ", ".join(have)
            if "source_original" not in cols:
                conn.execute(
                    f"""
                    INSERT INTO gbl_cs_lines_new ({insert_cols}, source_original)
                    SELECT {select_cols}, COALESCE(source_file, '')
                    FROM gbl_cs_lines
                    """
                )
            else:
                conn.execute(
                    f"""
                    INSERT INTO gbl_cs_lines_new ({insert_cols}, source_original)
                    SELECT {select_cols}, source_original FROM gbl_cs_lines
                    """
                )
            conn.execute("DROP TABLE gbl_cs_lines")
            conn.execute("ALTER TABLE gbl_cs_lines_new RENAME TO gbl_cs_lines")
        conn.execute(
            "CREATE INDEX IF NOT EXISTS idx_gbl_cs_month ON gbl_cs_lines(sail_month)"
        )
        conn.execute(
            "CREATE INDEX IF NOT EXISTS idx_gbl_cs_po ON gbl_cs_lines(po_no)"
        )
        conn.execute(
            "CREATE INDEX IF NOT EXISTS idx_gbl_cs_invoice ON gbl_cs_lines(invoice_no)"
        )
        conn.execute(
            "CREATE INDEX IF NOT EXISTS idx_gbl_cs_source_orig ON gbl_cs_lines(source_original)"
        )
        conn.commit()


def _normalize_invoice(invoice: str) -> str:
    return re.sub(r"\s+", "", (invoice or "").strip().upper())


def _workbook_key(name: str) -> str:
    """Stable key for a workbook so file.xlsx and file_2.xlsx replace each other.

    Normalize whitespace before stripping a trailing ``_N`` upload suffix so
    names like ``OCTOBER_2025 .xlsx`` (space before extension) still match.
    """
    raw = (name or "").strip()
    if not raw:
        return ""
    # Prefer stem when a path/filename is given; bare keys (no suffix) keep as-is
    stem = Path(raw).stem if ("." in Path(raw).name and Path(raw).suffix) else raw
    stem = re.sub(r"\s+", " ", stem).strip()
    # Only strip trailing _digits upload copies (file_2), not year tokens mid-name.
    # Apply after whitespace normalize so "FOO_2025 " → "FOO_2025" then year stays.
    # Upload copies look like stem_2 where the whole stem already ends with _N;
    # keep years (4-digit) — only strip 1–2 digit suffixes used by save_cs_upload.
    stem = re.sub(r"_(\d{1,2})$", "", stem)
    return stem.upper()


def list_input_files() -> list[dict[str, Any]]:
    """Uploaded workbooks currently contributing rows to the register."""
    init_register()
    with _connect() as conn:
        rows = conn.execute(
            """
            SELECT source_original, source_file, source_sheet, uploaded_at
            FROM gbl_cs_lines
            ORDER BY uploaded_at DESC, id DESC
            """
        ).fetchall()

    groups: dict[str, dict[str, Any]] = {}
    for r in rows:
        original = r["source_original"] or r["source_file"] or ""
        key = _workbook_key(original)
        if not key:
            continue
        g = groups.get(key)
        if not g:
            groups[key] = {
                "key": key,
                "display_name": original or r["source_file"] or key,
                "row_count": 1,
                "sheets": set(),
                "uploaded_at": r["uploaded_at"] or "",
                "saved_names": set(),
            }
            g = groups[key]
        else:
            g["row_count"] += 1
            if (r["uploaded_at"] or "") > (g["uploaded_at"] or ""):
                g["uploaded_at"] = r["uploaded_at"] or ""
                if original:
                    g["display_name"] = original
        if r["source_sheet"]:
            g["sheets"].add(r["source_sheet"])
        if r["source_file"]:
            g["saved_names"].add(r["source_file"])
        if original:
            g["saved_names"].add(Path(original).name)

    # Attach on-disk upload files that match each workbook key
    upload_dir = cs_upload_dir()
    disk_files = list(upload_dir.glob("*.xlsx")) + list(upload_dir.glob("*.xlsm"))
    for g in groups.values():
        matched = [
            p.name for p in disk_files if _workbook_key(p.name) == g["key"]
        ]
        g["disk_files"] = sorted(set(matched) | set(g.pop("saved_names")))
        g["sheets"] = ", ".join(sorted(g["sheets"])) if g["sheets"] else "—"

    return sorted(
        groups.values(),
        key=lambda x: x.get("uploaded_at") or "",
        reverse=True,
    )


def delete_input_workbook(key: str) -> dict[str, int]:
    """Remove register rows and saved upload files for a workbook key."""
    init_register()
    key = _workbook_key(key)
    if not key:
        raise ValueError("Missing file key.")

    with _connect() as conn:
        existing = conn.execute(
            "SELECT id, source_original, source_file FROM gbl_cs_lines"
        ).fetchall()
        drop_ids = [
            r["id"]
            for r in existing
            if _workbook_key(r["source_original"] or r["source_file"] or "") == key
        ]
        deleted_rows = 0
        if drop_ids:
            placeholders = ",".join("?" * len(drop_ids))
            deleted_rows = conn.execute(
                f"DELETE FROM gbl_cs_lines WHERE id IN ({placeholders})", drop_ids
            ).rowcount
        conn.commit()

    deleted_files = 0
    for path in cs_upload_dir().iterdir():
        if not path.is_file():
            continue
        if path.suffix.lower() not in {".xlsx", ".xlsm"}:
            continue
        if _workbook_key(path.name) == key:
            path.unlink(missing_ok=True)
            deleted_files += 1

    return {"deleted_rows": deleted_rows, "deleted_files": deleted_files}


def delete_register_row(row_id: int) -> bool:
    """Delete a single register line by id. Returns True if a row was removed."""
    init_register()
    with _connect() as conn:
        cur = conn.execute("DELETE FROM gbl_cs_lines WHERE id = ?", (int(row_id),))
        conn.commit()
        return cur.rowcount > 0


def merge_rows_into_register(
    rows: list[dict[str, Any]],
    *,
    source_file: str,
    source_original: str | None = None,
) -> dict[str, int]:
    """Replace this workbook's rows in the register, then insert the new set.

    Same invoice may exist from different workbooks; duplicates are flagged on display.
    Re-uploading the same original filename (or a _2 collision copy) replaces its lines.
    """
    init_register()
    now = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    original = (source_original or source_file or "").strip() or source_file
    key = _workbook_key(original)
    inserted = skipped = 0

    with _connect() as conn:
        existing = conn.execute(
            "SELECT id, source_original, source_file FROM gbl_cs_lines"
        ).fetchall()
        drop_ids = [
            r["id"]
            for r in existing
            if _workbook_key(r["source_original"] or r["source_file"] or "") == key
        ]
        deleted = 0
        if drop_ids:
            placeholders = ",".join("?" * len(drop_ids))
            deleted = conn.execute(
                f"DELETE FROM gbl_cs_lines WHERE id IN ({placeholders})", drop_ids
            ).rowcount
        for row in rows:
            invoice = (row.get("invoice_no") or "").strip()
            if not invoice:
                skipped += 1
                continue
            vals = {
                "invoice_no": invoice,
                "sr_no": row.get("sr_no") or "",
                "sail_month": (row.get("sail_month") or "").upper(),
                "vessel_sail_date": row.get("vessel_sail_date") or "",
                "po_no": row.get("po_no") or "",
                "po_date": row.get("po_date") or "",
                "ship_to_party": row.get("ship_to_party") or "",
                "port_of_discharge": row.get("port_of_discharge") or "",
                "qty": row.get("qty"),
                "currency": row.get("currency") or "",
                "rate_per_mt": row.get("rate_per_mt"),
                "value": row.get("value"),
                "ocean_freight": row.get("ocean_freight"),
                "insurance": row.get("insurance"),
                "fob_value": row.get("fob_value"),
                "commission": row.get("commission"),
                "commission_currency": row.get("commission_currency") or "",
                "source_file": source_file,
                "source_original": original,
                "source_sheet": row.get("source_sheet") or "",
                "product_hint": row.get("product_hint") or "",
                "uploaded_at": now,
            }
            conn.execute(
                """
                INSERT INTO gbl_cs_lines (
                  invoice_no, sr_no, sail_month, vessel_sail_date, po_no, po_date,
                  ship_to_party, port_of_discharge, qty, currency, rate_per_mt,
                  value, ocean_freight, insurance, fob_value, commission,
                  commission_currency, source_file, source_original, source_sheet,
                  product_hint, uploaded_at
                ) VALUES (
                  :invoice_no, :sr_no, :sail_month, :vessel_sail_date, :po_no, :po_date,
                  :ship_to_party, :port_of_discharge, :qty, :currency, :rate_per_mt,
                  :value, :ocean_freight, :insurance, :fob_value, :commission,
                  :commission_currency, :source_file, :source_original, :source_sheet,
                  :product_hint, :uploaded_at
                )
                """,
                vals,
            )
            inserted += 1
        conn.commit()

    return {
        "inserted": inserted,
        "replaced": max(deleted, 0),
        "skipped": skipped,
    }


def list_register_rows(
    *,
    month: str | None = None,
    po: str | None = None,
) -> list[dict[str, Any]]:
    init_register()
    clauses: list[str] = []
    params: list[Any] = []
    if month:
        clauses.append("UPPER(sail_month) = ?")
        params.append(month.strip().upper())
    if po:
        clauses.append("po_no LIKE ?")
        params.append(f"%{po.strip()}%")
    where = ("WHERE " + " AND ".join(clauses)) if clauses else ""
    sql = f"""
        SELECT * FROM gbl_cs_lines
        {where}
        ORDER BY
          CASE UPPER(sail_month)
            WHEN 'APRIL' THEN 4 WHEN 'MAY' THEN 5 WHEN 'JUNE' THEN 6
            WHEN 'JULY' THEN 7 WHEN 'AUGUST' THEN 8 WHEN 'SEPTEMBER' THEN 9
            WHEN 'OCTOBER' THEN 10 WHEN 'NOVEMBER' THEN 11 WHEN 'DECEMBER' THEN 12
            WHEN 'JANUARY' THEN 13 WHEN 'FEBRUARY' THEN 14 WHEN 'MARCH' THEN 15
            ELSE 99
          END,
          invoice_no,
          id
    """
    with _connect() as conn:
        rows = [dict(r) for r in conn.execute(sql, params).fetchall()]
        # Duplicate counts are by exact invoice # (normalized), across the full register.
        # AA vs IAA are different keys — each is only a duplicate of itself.
        counts: dict[str, int] = {}
        occur_months: dict[str, list[str]] = {}
        for r in conn.execute(
            "SELECT invoice_no, sail_month, source_original FROM gbl_cs_lines"
        ).fetchall():
            key = _normalize_invoice(r["invoice_no"])
            if not key:
                continue
            counts[key] = counts.get(key, 0) + 1
            month_label = (r["sail_month"] or "").strip().upper() or "?"
            src = (r["source_original"] or "").strip()
            tip = f"{month_label}" + (f" ({src})" if src else "")
            occur_months.setdefault(key, []).append(tip)

    for row in rows:
        key = _normalize_invoice(row.get("invoice_no") or "")
        n = counts.get(key, 0)
        row["duplicate_count"] = n
        row["is_duplicate"] = n > 1
        if n > 1:
            places = occur_months.get(key) or []
            # Dedupe while preserving order
            seen: set[str] = set()
            uniq_places: list[str] = []
            for p in places:
                if p not in seen:
                    seen.add(p)
                    uniq_places.append(p)
            inv = (row.get("invoice_no") or "").strip()
            row["duplicate_flag"] = f"DUPLICATE ({n})"
            row["duplicate_detail"] = (
                f"Exact invoice {inv} appears {n} times: " + "; ".join(uniq_places)
            )
        else:
            row["duplicate_flag"] = ""
            row["duplicate_detail"] = ""
    return rows


def list_register_rows_for_workbook_keys(keys: list[str]) -> list[dict[str, Any]]:
    """Register lines belonging to any of the given workbook keys."""
    wanted = {_workbook_key(k) for k in (keys or []) if (k or "").strip()}
    if not wanted:
        return []
    out = []
    for row in list_register_rows():
        src = row.get("source_original") or row.get("source_file") or ""
        if _workbook_key(src) in wanted:
            out.append(row)
    return out


def register_filter_options() -> dict[str, list[str]]:
    init_register()
    with _connect() as conn:
        months = [
            r[0]
            for r in conn.execute(
                """
                SELECT DISTINCT UPPER(sail_month) FROM gbl_cs_lines
                WHERE sail_month IS NOT NULL AND TRIM(sail_month) != ''
                ORDER BY
                  CASE UPPER(sail_month)
                    WHEN 'APRIL' THEN 4 WHEN 'MAY' THEN 5 WHEN 'JUNE' THEN 6
                    WHEN 'JULY' THEN 7 WHEN 'AUGUST' THEN 8 WHEN 'SEPTEMBER' THEN 9
                    WHEN 'OCTOBER' THEN 10 WHEN 'NOVEMBER' THEN 11 WHEN 'DECEMBER' THEN 12
                    WHEN 'JANUARY' THEN 13 WHEN 'FEBRUARY' THEN 14 WHEN 'MARCH' THEN 15
                    ELSE 99
                  END
                """
            ).fetchall()
        ]
        po_raw = [
            r[0]
            for r in conn.execute(
                """
                SELECT DISTINCT po_no FROM gbl_cs_lines
                WHERE po_no IS NOT NULL AND TRIM(po_no) != ''
                ORDER BY po_no
                """
            ).fetchall()
        ]

    po_numbers: set[str] = set()
    for blob in po_raw:
        for part in re.split(r"[,;\n]+", blob):
            p = part.strip()
            if p:
                po_numbers.add(p)
    return {"months": months, "po_numbers": sorted(po_numbers)}


def register_stats(
    *,
    month: str | None = None,
    po: str | None = None,
) -> dict[str, Any]:
    init_register()
    clauses: list[str] = []
    params: list[Any] = []
    if month:
        clauses.append("UPPER(sail_month) = ?")
        params.append(month.strip().upper())
    if po:
        clauses.append("po_no LIKE ?")
        params.append(f"%{po.strip()}%")
    where = ("WHERE " + " AND ".join(clauses)) if clauses else ""
    with _connect() as conn:
        total = conn.execute(
            f"SELECT COUNT(*) FROM gbl_cs_lines {where}", params
        ).fetchone()[0]
        files = conn.execute(
            f"SELECT COUNT(DISTINCT source_file) FROM gbl_cs_lines {where}", params
        ).fetchone()[0]
        qty = conn.execute(
            f"SELECT COALESCE(SUM(qty), 0) FROM gbl_cs_lines {where}", params
        ).fetchone()[0]
        sale = conn.execute(
            f"SELECT COALESCE(SUM(value), 0) FROM gbl_cs_lines {where}", params
        ).fetchone()[0]
        commission = conn.execute(
            f"SELECT COALESCE(SUM(commission), 0) FROM gbl_cs_lines {where}", params
        ).fetchone()[0]
    return {
        "row_count": total,
        "file_count": files,
        "qty_total": qty,
        "sale_total": sale,
        "commission_total": commission,
    }


def recompute_sail_months_from_dates() -> int:
    """Backfill sail_month on existing register rows from vessel_sail_date."""
    init_register()
    updated = 0
    with _connect() as conn:
        rows = conn.execute(
            "SELECT id, vessel_sail_date, sail_month FROM gbl_cs_lines"
        ).fetchall()
        for row in rows:
            month = _month_from_sail_date(row["vessel_sail_date"])
            if not month or month == (row["sail_month"] or "").upper():
                continue
            conn.execute(
                "UPDATE gbl_cs_lines SET sail_month = ? WHERE id = ?",
                (month, row["id"]),
            )
            updated += 1
        conn.commit()
    return updated
